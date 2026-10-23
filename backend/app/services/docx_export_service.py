from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

from backend.app.core.config import settings
from backend.app.schemas.brd import BRDDocument


class DocxExportService:
    @staticmethod
    def generate_docx(brd: BRDDocument) -> Path:
        doc = Document()

        # Document Title / Cover Header
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(36)
        title_p.paragraph_format.space_after = Pt(12)
        title_run = title_p.add_run(brd.project_title)
        title_run.font.size = Pt(26)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(24, 43, 73)

        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run("Business Requirements Document (BRD) — Reverse-Engineered from Source")
        sub_run.font.size = Pt(14)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(90, 105, 120)

        # Meta Table
        table = doc.add_table(rows=5, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_data = [
            ("Document ID", brd.id),
            ("Version", f"{brd.version}.0"),
            ("Analysis ID", brd.analysis_id),
            ("Target Framework", brd.target_framework),
            ("Date Generated", brd.generated_at),
        ]
        for row_idx, (k, v) in enumerate(meta_data):
            r = table.rows[row_idx]
            r.cells[0].paragraphs[0].add_run(k).bold = True
            r.cells[1].paragraphs[0].add_run(str(v))

        doc.add_page_break()

        # Section 1: Executive Summary
        h1 = doc.add_heading("1. Executive Summary", level=1)
        doc.add_paragraph(brd.executive_summary)

        # Section 2: Business Objectives
        doc.add_heading("2. Business Objectives", level=1)
        for obj in brd.business_objectives:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(obj)

        # Section 3: Discovered Modules & Architecture
        doc.add_heading("3. Discovered Modules & Scope", level=1)
        doc.add_paragraph(f"Scope: {brd.application_scope}")
        mod_table = doc.add_table(rows=1, cols=3)
        mod_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = mod_table.rows[0].cells
        hdr_cells[0].paragraphs[0].add_run("Module Name").bold = True
        hdr_cells[1].paragraphs[0].add_run("Source Directory").bold = True
        hdr_cells[2].paragraphs[0].add_run("Key Responsibilities").bold = True

        for mod in brd.modules:
            row = mod_table.add_row().cells
            row[0].paragraphs[0].add_run(mod.name)
            row[1].paragraphs[0].add_run(mod.source_directory)
            row[2].paragraphs[0].add_run(", ".join(mod.responsibilities))

        # Section 4: Actors & Security Roles
        doc.add_heading("4. Actors and Security Roles", level=1)
        role_table = doc.add_table(rows=1, cols=3)
        role_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        r_hdr = role_table.rows[0].cells
        r_hdr[0].paragraphs[0].add_run("Role").bold = True
        r_hdr[1].paragraphs[0].add_run("Permissions").bold = True
        r_hdr[2].paragraphs[0].add_run("Code Reference").bold = True

        for act in brd.actors_and_roles:
            r = role_table.add_row().cells
            r[0].paragraphs[0].add_run(act.role_name)
            r[1].paragraphs[0].add_run(", ".join(act.discovered_permissions))
            r[2].paragraphs[0].add_run(act.source_reference or "N/A")

        # Section 5: Functional Requirements & Evidence Traceability
        doc.add_heading("5. Functional Requirements & Code Traceability", level=1)
        for req in brd.functional_requirements:
            doc.add_heading(f"{req.id}: {req.title}", level=2)
            p_desc = doc.add_paragraph()
            p_desc.add_run("Description: ").bold = True
            p_desc.add_run(req.description)

            p_meta = doc.add_paragraph()
            p_meta.add_run(f"Category: {req.category} | Confidence: {req.confidence.value} | Classification: {req.classification.value}")

            p_ev = doc.add_paragraph()
            p_ev.add_run("Supporting Evidence IDs: ").bold = True
            p_ev.add_run(", ".join(req.supporting_evidence_ids))

            if req.acceptance_criteria:
                doc.add_heading("Acceptance Criteria (Gherkin Scenarios):", level=3)
                for ac in req.acceptance_criteria:
                    ac_p = doc.add_paragraph(style='List Bullet')
                    ac_p.add_run(f"[{ac.id}] {ac.scenario}\n").bold = True
                    ac_p.add_run(f"  Given: {ac.given}\n")
                    ac_p.add_run(f"  When: {ac.when}\n")
                    ac_p.add_run(f"  Then: {ac.then}\n")

        # Section 6: Business Rules & Validation Thresholds
        doc.add_heading("6. Business Rules & Thresholds", level=1)
        br_table = doc.add_table(rows=1, cols=4)
        br_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        b_hdr = br_table.rows[0].cells
        b_hdr[0].paragraphs[0].add_run("Rule ID").bold = True
        b_hdr[1].paragraphs[0].add_run("Rule Title").bold = True
        b_hdr[2].paragraphs[0].add_run("Threshold").bold = True
        b_hdr[3].paragraphs[0].add_run("Source Symbol").bold = True

        for br in brd.business_rules:
            row = br_table.add_row().cells
            row[0].paragraphs[0].add_run(br.id)
            row[1].paragraphs[0].add_run(br.title)
            th_str = f"{br.threshold_operator or ''} {br.threshold_value or ''}".strip() or "Logic check"
            row[2].paragraphs[0].add_run(th_str)
            row[3].paragraphs[0].add_run(br.source_symbol)

        # Section 7: Open Questions & Stakeholder Action Items
        doc.add_heading("7. Stakeholder Open Questions", level=1)
        for oq in brd.open_questions:
            doc.add_heading(f"{oq.id}: {oq.topic} (Target: {oq.stakeholder_target})", level=2)
            p = doc.add_paragraph()
            p.add_run(f"Question: {oq.question}\n")
            p.add_run(f"Code Anchor: {oq.context_code_reference} | Priority: {oq.urgency}")

        # Save docx
        export_filename = f"BRD_{brd.analysis_id[:8]}_{datetime.now().strftime('%Y%m%d%H%M%S')}.docx"
        out_path = settings.EXPORTS_DIR / export_filename
        doc.save(str(out_path))
        return out_path
