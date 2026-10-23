from pathlib import Path

from docx import Document

from backend.app.agents.bedrock_provider import DeterministicDemoProvider
from backend.app.code_intelligence.csharp_scanner import CSharpScanner
from backend.app.services.docx_export_service import DocxExportService


def test_docx_export_generates_valid_document():
    sample_dir = Path(__file__).resolve().parent.parent / "samples" / "LegacyOrders"
    scanner = CSharpScanner("test_repo_docx", sample_dir)
    summary, symbols, evidence, _ = scanner.scan_all()

    provider = DeterministicDemoProvider()
    brd = provider.generate_brd("test_analysis_docx", summary, symbols, evidence)

    out_file = DocxExportService.generate_docx(brd)
    assert out_file.exists()
    assert out_file.stat().st_size > 1000

    # Verify docx readability
    doc = Document(str(out_file))
    paragraphs_text = [p.text for p in doc.paragraphs]
    assert any("Business Requirements Document" in p for p in paragraphs_text)
    assert any("Executive Summary" in p for p in paragraphs_text)
    assert len(doc.tables) >= 3
